#!/usr/bin/env python3
"""Aggiunge i campi micronutrienti a bls_2025.json (gia' esistente),
leggendo il foglio dati ufficiale BLS 4.0 (Bundeslebensmittelschlüssel,
Max Rubner-Institut). Vedi DECISIONI.md "Nota backlog: micronutrienti".

Fonte: BLS_4_0_Daten_2025_DE.xlsx (foglio dati, formato largo: 3 colonne
per nutriente "CODICE nome [unita]" / "CODICE Datenherkunft" / "CODICE
Referenz"), colonna "BLS Code" = food_code gia' presente in bls_2025.json.
Codici nutriente identificati da BLS_4_0_Components_DE_EN.xlsx (dizionario
componenti, non serve a runtime: i codici sono gia' hardcoded sotto).

Note sui valori:
- Cella vuota/None -> 0.0, stesso principio "non disponibile" gia' usato
  per CREA/CIQUAL.
- **Selenio sempre 0.0**: BLS non lo include come nutriente tracciato
  (verificato cercando "selen" sia in tedesco che in inglese nel
  dizionario componenti, nessun codice trovato) — limite della fonte,
  diverso da CREA (che ce l'ha) e da CIQUAL (che ce l'ha).
- Vitamina A: usato VITA (retinol equivalents), non VITAA (RAE) — stessa
  granularita' "retinol equivalent" di CIQUAL/CREA.
- Vitamina K: BLS ha gia' un codice aggregato VITK (a differenza di
  CIQUAL che va sommato da K1+K2) — usato direttamente.
- Folati: usato FOL (folate equivalent, aggregato), non FOLFD (folato
  grezzo) ne' FOLAC (acido folico sintetico) — stessa granularita'
  "folati totali" delle altre due fonti.
- **Conversione unita' µg->mg per rame/manganese/vitamina B6**: a
  differenza di CIQUAL/CREA (entrambe in mg/100g per questi 3 campi),
  il foglio dati BLS li riporta in µg/100g (verificato negli header:
  "CU Kupfer [µg/100g]", "MN Mangan [µg/100g]", "VITB6 Vitamin B6
  [µg/100g]") — senza conversione i valori sarebbero risultati 1000
  volte troppo alti rispetto alle altre due fonti (es. rame ~180 invece
  di ~0.18). Tutti gli altri campi sono gia' nella stessa unita' delle
  altre due fonti (verificato uno per uno negli header del foglio dati),
  nessun'altra conversione necessaria.
"""
import json

DATI_XLSX = "BLS_4_0_Daten_2025_DE.xlsx"
BLS_JSON_PATH = "bls_2025.json"

# campo_output -> codice nutriente BLS (None = non disponibile nella fonte)
CAMPI_MICRONUTRIENTI = {
    "sodio": "NA",
    "potassio": "K",
    "calcio": "CA",
    "ferro": "FE",
    "magnesio": "MG",
    "zinco": "ZN",
    "fosforo": "P",
    "iodio": "ID",
    "selenio": None,  # non presente nella fonte BLS
    "rame": "CU",
    "manganese": "MN",
    "vitamina_a": "VITA",
    "vitamina_b1": "THIA",
    "vitamina_b2": "RIBF",
    "vitamina_b3": "NIA",
    "vitamina_b5": "PANTAC",
    "vitamina_b6": "VITB6",
    "vitamina_b7": "BIOT",
    "vitamina_b9": "FOL",
    "vitamina_b12": "VITB12",
    "vitamina_c": "VITC",
    "vitamina_d": "VITD",
    "vitamina_e": "VITE",
    "vitamina_k": "VITK",
    "colesterolo": "CHORL",
    "grassi_saturi": "FASAT",
    "grassi_monoinsaturi": "FAMS",
    "grassi_polinsaturi": "FAPU",
    "omega3_ala": "F18:3CN3",
    "omega3_epa": "F20:5CN3",
    "omega3_dha": "F22:6CN3",
    "omega6_linoleico": "F18:2CN6",
}

# Campi il cui valore nel foglio dati BLS e' in µg/100g mentre le altre
# due fonti (CIQUAL/CREA) usano mg/100g per lo stesso campo — convertiti
# dividendo per 1000. Vedi nota nel docstring sopra.
CAMPI_DA_CONVERTIRE_UG_A_MG = {"rame", "manganese", "vitamina_b6"}


def to_float(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def main() -> None:
    import openpyxl

    wb = openpyxl.load_workbook(DATI_XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)

    colonna_per_codice = {}
    for campo, codice in CAMPI_MICRONUTRIENTI.items():
        if codice is None:
            continue
        prefisso = f"{codice} "
        trovata = None
        for i, h in enumerate(header):
            if h and str(h).startswith(prefisso):
                trovata = i
                break
        colonna_per_codice[codice] = trovata

    valori_per_food_code = {}
    for row in rows:
        if not row or not row[0]:
            continue
        food_code = str(row[0]).strip()
        valori = {}
        for campo, codice in CAMPI_MICRONUTRIENTI.items():
            if codice is None:
                valori[campo] = 0.0
                continue
            idx = colonna_per_codice.get(codice)
            valore = to_float(row[idx]) if idx is not None else 0.0
            if campo in CAMPI_DA_CONVERTIRE_UG_A_MG:
                valore /= 1000.0
            valori[campo] = valore
        valori_per_food_code[food_code] = valori

    with open(BLS_JSON_PATH, encoding="utf-8") as f:
        alimenti = json.load(f)

    trovati = 0
    for alimento in alimenti:
        valori = valori_per_food_code.get(alimento["food_code"])
        if valori is None:
            for campo in CAMPI_MICRONUTRIENTI:
                alimento[campo] = 0.0
            continue
        trovati += 1
        alimento.update(valori)

    with open(BLS_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(alimenti, f, ensure_ascii=False, indent=None, separators=(",", ":"))

    print(f"Alimenti totali: {len(alimenti)}")
    print(f"Con food_code trovato nel foglio dati: {trovati}")
    print(f"Senza match: {len(alimenti) - trovati}")
    print("Colonne non trovate per codice:", [c for c, i in colonna_per_codice.items() if i is None])


if __name__ == "__main__":
    main()
